from flask import Blueprint


import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl import load_workbook
import re
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.exam_batch_service import (
    EXAM_STUDENT_STATUS_LABELS,
    classify_exam_student_status,
)
from services.exam_paper_service import ensure_exam_paper_enabled


def create_admin_blueprint(deps):
    globals().update(deps)
    bp = Blueprint('admin', __name__)

    @bp.route('/admin')
    @login_required
    def admin_dashboard():
        """管理员主页"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
    
        selected_paper_id = resolve_selected_exam_paper_id(
            request.args.get('paper_id', type=int),
            include_disabled_for_admin=True
        )
        selected_paper = get_exam_paper_by_id(selected_paper_id) if selected_paper_id else None
        total_problems = get_total_problem_count(selected_paper_id)

        return render_template(
            'admin_dashboard.html',
            exam_papers=get_exam_papers(include_disabled=True),
            selected_paper_id=selected_paper_id,
            selected_paper=selected_paper,
            stats={'stats': {}, 'today_stats': {}},
            total_problems=total_problems,
            recent_completions=[],
            incomplete_students=[],
            exams=get_all_exams(),
        )



    def _get_exam_setup_options():
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("""
                SELECT DISTINCT TRIM(class_name) AS value
                FROM users
                WHERE username != 'admin' AND class_name IS NOT NULL AND TRIM(class_name) != ''
                ORDER BY value
            """)
            class_options = [row['value'] for row in cursor.fetchall()]
            cursor.execute("""
                SELECT DISTINCT TRIM(major) AS value
                FROM users
                WHERE username != 'admin' AND major IS NOT NULL AND TRIM(major) != ''
                ORDER BY value
            """)
            major_options = [row['value'] for row in cursor.fetchall()]
            cursor.execute("""
                SELECT DISTINCT TRIM(teacher_name) AS value
                FROM users
                WHERE username != 'admin' AND teacher_name IS NOT NULL AND TRIM(teacher_name) != ''
                ORDER BY value
            """)
            course_options = [row['value'] for row in cursor.fetchall()]
            return class_options, major_options, course_options
        finally:
            cursor.close()
            conn.close()


    def _finish_auto_enabled_paper(paper_id, auto_enabled):
        if not auto_enabled:
            return
        invalidate_exam_paper_cache(paper_id)
        schedule_exam_paper_prewarm(paper_id)
        flash('所选题库原本处于关闭状态，已随考试设置自动开启并开始预热题目池。', 'info')


    def _normalize_text(value):
        return str(value).strip() if value is not None else ''


    def _clean_student_id(value):
        text = _normalize_text(value)
        if not text:
            return ''
        text = re.sub(r'\s+', '', text)
        if re.fullmatch(r'\d+\.0', text):
            text = text[:-2]
        if 'e' in text.lower():
            try:
                text = str(int(float(text)))
            except Exception:
                return ''
        text = re.sub(r'\D', '', text)
        return text


    def _match_header_index(header_row, candidates):
        normalized = {_normalize_text(item).lower() for item in candidates}
        for index, value in enumerate(header_row):
            if _normalize_text(value).lower() in normalized:
                return index
        return None


    def _build_missing_accounts_workbook(rows, exam_name):
        wb = Workbook()
        ws = wb.active
        ws.title = '未找到名单'
        ws.append(['行号', '学号/姓名', '姓名', '原因'])
        for item in rows:
            ws.append([
                item.get('row_no'),
                item.get('identifier') or '',
                item.get('name') or '',
                item.get('reason') or '',
            ])

        if rows:
            for col in ws.columns:
                column_letter = get_column_letter(col[0].column)
                max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col[:200])
                ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.freeze_panes = 'A2'

        meta = wb.create_sheet('说明')
        meta['A1'] = '考试名称'
        meta['B1'] = exam_name
        meta['A2'] = '说明'
        meta['B2'] = '仅列出未匹配到系统账户的 Excel 行。'
        meta.column_dimensions['A'].width = 16
        meta.column_dimensions['B'].width = 60
        return wb


    def _build_missing_accounts_response(rows, exam_name):
        workbook = _build_missing_accounts_workbook(rows, exam_name)
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        safe_name = re.sub(r'[^A-Za-z0-9_-]+', '_', exam_name or 'retake_missing_accounts').strip('_') or 'retake_missing_accounts'
        filename = f"{safe_name}_missing_accounts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response


    def _parse_exam_import_rows(upload_file):
        workbook = load_workbook(upload_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header = rows[0]
        id_idx = _match_header_index(header, {'学号', '学生学号', 'student_id', 'studentid', 'id', '账号', '用户名', 'username'})
        name_idx = _match_header_index(header, {'姓名', '名字', 'name', 'student_name'})

        has_header = id_idx is not None or name_idx is not None
        start_row = 2 if has_header else 1
        if id_idx is None and name_idx is None:
            id_idx = 0
            name_idx = 1 if len(header) > 1 else None

        parsed_rows = []
        for excel_index, row in enumerate(rows[start_row - 1:], start=start_row):
            student_id = _clean_student_id(row[id_idx] if id_idx is not None and id_idx < len(row) else None)
            student_name = _normalize_text(row[name_idx]) if name_idx is not None and name_idx < len(row) else ''
            if not student_id and not student_name:
                continue
            parsed_rows.append({
                'row_no': excel_index,
                'student_id': student_id,
                'student_name': student_name,
                'identifier': student_id or student_name,
            })
        return parsed_rows


    def _resolve_assignment_user_ids(cursor, assignments):
        user_ids = set()
        field_map = {
            'class': 'class_name',
            'major': 'major',
            'course': 'teacher_name',
        }
        for assign_type, assign_value in assignments:
            if assign_type == 'student':
                try:
                    user_ids.add(int(assign_value))
                except (TypeError, ValueError):
                    continue
                continue
            field_name = field_map.get(assign_type)
            if not field_name:
                continue
            cursor.execute(
                f"SELECT id FROM users WHERE username != 'admin' AND BINARY COALESCE({field_name}, '') = BINARY %s",
                (assign_value,),
            )
            user_ids.update(int(row['id']) for row in cursor.fetchall())
        return sorted(user_ids)


    def _resolve_uploaded_student_assignments(cursor, upload_file):
        parsed_rows = _parse_exam_import_rows(upload_file)
        matched_ids = set()
        missing_count = 0
        for row in parsed_rows:
            matched_user = None
            if row['student_id']:
                cursor.execute(
                    "SELECT id FROM users WHERE username != 'admin' AND username = %s LIMIT 1",
                    (row['student_id'],),
                )
                matched_user = cursor.fetchone()
            if not matched_user and row['student_name']:
                cursor.execute(
                    "SELECT id FROM users WHERE username != 'admin' AND TRIM(name) = %s LIMIT 2",
                    (row['student_name'],),
                )
                matches = cursor.fetchall()
                matched_user = matches[0] if len(matches) == 1 else None
            if matched_user:
                matched_ids.add(int(matched_user['id']))
            else:
                missing_count += 1
        return [('student', str(user_id)) for user_id in sorted(matched_ids)], missing_count


    def _create_exam_batch(cursor, exam_id, name, start_time, end_time, assignments, is_default=False):
        batch_name = (name or '').strip() or ('默认批次' if is_default else '考试批次')
        user_ids = _resolve_assignment_user_ids(cursor, assignments)
        if not user_ids:
            raise ValueError('当前分配范围内没有可用学生。')

        placeholders = ', '.join(['%s'] * len(user_ids))
        cursor.execute(
            f"""
            SELECT u.username, u.name, b.name AS batch_name
            FROM exam_batch_students s
            JOIN users u ON u.id = s.user_id
            JOIN exam_batches b ON b.id = s.batch_id
            WHERE s.exam_id = %s AND s.user_id IN ({placeholders})
            LIMIT 5
            """,
            [exam_id, *user_ids],
        )
        conflicts = cursor.fetchall()
        if conflicts:
            names = '、'.join((row.get('name') or row.get('username') or '') for row in conflicts)
            raise ValueError(f'以下学生已属于本考试的其他批次：{names}')

        cursor.execute(
            """
            INSERT INTO exam_batches (exam_id, name, start_time, end_time, status, is_default)
            VALUES (%s, %s, %s, %s, 'active', %s)
            """,
            (exam_id, batch_name, start_time, end_time, bool(is_default)),
        )
        batch_id = cursor.lastrowid
        cursor.executemany(
            """
            INSERT INTO exam_batch_assignments (batch_id, assign_type, assign_value)
            VALUES (%s, %s, %s)
            """,
            [(batch_id, assign_type, assign_value) for assign_type, assign_value in assignments],
        )
        cursor.executemany(
            """
            INSERT INTO exam_batch_students (exam_id, batch_id, user_id)
            VALUES (%s, %s, %s)
            """,
            [(exam_id, batch_id, user_id) for user_id in user_ids],
        )
        return batch_id, len(user_ids)


    def _replace_exam_batch_assignments(cursor, exam_id, batch_id, assignments):
        user_ids = _resolve_assignment_user_ids(cursor, assignments)
        if not user_ids:
            raise ValueError('当前分配范围内没有可用学生。')
        placeholders = ', '.join(['%s'] * len(user_ids))
        cursor.execute(
            f"""
            SELECT u.username, u.name, b.name AS batch_name
            FROM exam_batch_students s
            JOIN users u ON u.id = s.user_id
            JOIN exam_batches b ON b.id = s.batch_id
            WHERE s.exam_id = %s AND s.batch_id != %s
              AND s.user_id IN ({placeholders})
            LIMIT 5
            """,
            [exam_id, batch_id, *user_ids],
        )
        conflicts = cursor.fetchall()
        if conflicts:
            names = '、'.join((row.get('name') or row.get('username') or '') for row in conflicts)
            raise ValueError(f'以下学生已属于本考试的其他批次：{names}')
        cursor.execute('DELETE FROM exam_batch_students WHERE batch_id = %s', (batch_id,))
        cursor.execute('DELETE FROM exam_batch_assignments WHERE batch_id = %s', (batch_id,))
        cursor.executemany(
            "INSERT INTO exam_batch_assignments (batch_id, assign_type, assign_value) VALUES (%s, %s, %s)",
            [(batch_id, assign_type, assign_value) for assign_type, assign_value in assignments],
        )
        cursor.executemany(
            "INSERT INTO exam_batch_students (exam_id, batch_id, user_id) VALUES (%s, %s, %s)",
            [(exam_id, batch_id, user_id) for user_id in user_ids],
        )
        return len(user_ids)


    @bp.route('/admin/exams')
    @login_required
    def admin_exams():
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        class_options, major_options, course_options = _get_exam_setup_options()
        edit_exam_id = request.args.get('edit_exam_id', type=int)
        editing_exam = get_exam_by_id(edit_exam_id) if edit_exam_id else None
        if editing_exam and editing_exam.get('status') == 'archived':
            flash('归档考试只允许查看历史数据，不能再编辑。', 'warning')
            return redirect(url_for('admin.admin_exam_detail', exam_id=edit_exam_id))
        editing_assignments = []
        editing_assign_type = ''
        editing_assign_value = ''
        editing_assign_values = []
        if edit_exam_id:
            conn = get_db_connection()
            cursor = conn.cursor(dictionary=True)
            try:
                cursor.execute(
                    """
                    SELECT a.assign_type, a.assign_value,
                           b.id AS batch_id, b.name AS batch_name,
                           b.start_time, b.end_time
                    FROM exam_batches b
                    JOIN exam_batch_assignments a ON a.batch_id = b.id
                    WHERE b.exam_id = %s AND b.is_default = TRUE
                    ORDER BY a.id
                    """,
                    (edit_exam_id,)
                )
                editing_assignments = cursor.fetchall()
                if editing_assignments:
                    editing_exam['start_time'] = editing_assignments[0]['start_time']
                    editing_exam['end_time'] = editing_assignments[0]['end_time']
                    editing_exam['batch_id'] = editing_assignments[0]['batch_id']
                    editing_exam['batch_name'] = editing_assignments[0]['batch_name']
                    editing_assign_type = editing_assignments[0]['assign_type']
                    editing_assign_value = editing_assignments[0]['assign_value']
                    editing_assign_values = [
                        item['assign_value']
                        for item in editing_assignments
                        if item['assign_type'] == editing_assign_type
                    ]
            finally:
                cursor.close()
                conn.close()
        return render_template(
            'admin_exams.html',
            exams=get_all_exams(),
            exam_papers=get_exam_papers(include_disabled=True),
            class_options=class_options,
            major_options=major_options,
            course_options=course_options,
            retake_candidates=[],
            editing_exam=editing_exam,
            editing_assignments=editing_assignments,
            editing_assign_type=editing_assign_type,
            editing_assign_value=editing_assign_value,
            editing_assign_values=editing_assign_values,
            selected_source_exam_id=request.args.get('source_exam_id', type=int),
            format_datetime_local=format_datetime_local
        )


    @bp.route('/admin/exams/<int:exam_id>')
    @login_required
    def admin_exam_detail(exam_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        exam = get_exam_by_id(exam_id)
        if not exam:
            flash('考试不存在或已被删除。', 'danger')
            return redirect(url_for('admin.admin_exams'))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute(
                """
                SELECT b.*,
                       COUNT(DISTINCT s.user_id) AS student_count,
                       COUNT(DISTINCT q.user_id) AS entered_count
                FROM exam_batches b
                LEFT JOIN exam_batch_students s ON s.batch_id = b.id
                LEFT JOIN exam_user_questions q ON q.batch_id = b.id
                WHERE b.exam_id = %s
                GROUP BY b.id
                ORDER BY b.start_time, b.id
                """,
                (exam_id,),
            )
            batches = cursor.fetchall()
            selected_batch_id = request.args.get('batch_id', type=int)
            valid_batch_ids = {int(batch['id']) for batch in batches}
            if selected_batch_id not in valid_batch_ids:
                selected_batch_id = int(batches[0]['id']) if batches else None

            students = []
            if selected_batch_id:
                cursor.execute(
                    """
                    SELECT u.id, u.username, u.name, u.class_name, u.major, u.teacher_name,
                           COALESCE(q.entered_count, 0) AS entered_count,
                           COALESCE(c.completed_count, 0) AS completed_count,
                           COALESCE(r.answer_rows, 0) AS answer_rows,
                           COALESCE(r.correct_rows, 0) AS correct_rows,
                           r.first_response_time, r.last_response_time
                    FROM exam_batch_students s
                    JOIN users u ON u.id = s.user_id
                    LEFT JOIN (
                        SELECT user_id, COUNT(*) AS entered_count
                        FROM exam_user_questions
                        WHERE exam_id = %s AND batch_id = %s
                        GROUP BY user_id
                    ) q ON q.user_id = u.id
                    LEFT JOIN (
                        SELECT user_id, COUNT(DISTINCT template_id) AS completed_count
                        FROM (
                            SELECT user_id, template_id, attempt_count
                            FROM user_responses
                            WHERE exam_id = %s AND batch_id = %s
                            GROUP BY user_id, template_id, attempt_count
                            HAVING SUM(CASE WHEN is_correct THEN 1 ELSE 0 END) = COUNT(*)
                        ) completed_attempts
                        GROUP BY user_id
                    ) c ON c.user_id = u.id
                    LEFT JOIN (
                        SELECT user_id, COUNT(*) AS answer_rows,
                               SUM(CASE WHEN is_correct THEN 1 ELSE 0 END) AS correct_rows,
                               MIN(response_time) AS first_response_time,
                               MAX(response_time) AS last_response_time
                        FROM user_responses
                        WHERE exam_id = %s AND batch_id = %s
                        GROUP BY user_id
                    ) r ON r.user_id = u.id
                    WHERE s.exam_id = %s AND s.batch_id = %s
                    ORDER BY u.class_name, u.username
                    """,
                    (
                        exam_id, selected_batch_id,
                        exam_id, selected_batch_id,
                        exam_id, selected_batch_id,
                        exam_id, selected_batch_id,
                    ),
                )
                students = cursor.fetchall()
        finally:
            cursor.close()
            conn.close()

        selected_batch = next(
            (batch for batch in batches if int(batch['id']) == selected_batch_id),
            None,
        )
        total_problems = int(
            exam['question_count']
            if exam.get('question_count') is not None
            else get_exam_question_pool_count(exam_id, exam['paper_id'])
        )
        now = datetime.now()
        status_labels = EXAM_STUDENT_STATUS_LABELS
        status_counts = {key: 0 for key in status_labels}
        for student in students:
            entered = int(student.get('entered_count') or 0) > 0
            status = classify_exam_student_status(
                entered=entered,
                completed_count=int(student.get('completed_count') or 0),
                total_problems=total_problems,
                start_time=selected_batch.get('start_time') if selected_batch else None,
                end_time=selected_batch.get('end_time') if selected_batch else None,
                now=now,
            )
            student['exam_status'] = status
            student['exam_status_label'] = status_labels[status]
            student['accuracy'] = round(
                int(student.get('correct_rows') or 0) / int(student.get('answer_rows') or 1) * 100,
                1,
            ) if student.get('answer_rows') else None
            status_counts[status] += 1

        batch_class_options = sorted({
            str(student.get('class_name') or '').strip()
            for student in students
            if str(student.get('class_name') or '').strip()
        })
        batch_major_options = sorted({
            str(student.get('major') or '').strip()
            for student in students
            if str(student.get('major') or '').strip()
        })
        batch_course_options = sorted({
            str(student.get('teacher_name') or '').strip()
            for student in students
            if str(student.get('teacher_name') or '').strip()
        })

        selected_status = (request.args.get('status') or '').strip()
        selected_class_name = (request.args.get('class_name') or '').strip()
        selected_major = (request.args.get('major') or '').strip()
        selected_course = (request.args.get('course') or '').strip()
        visible_students = [
            student for student in students
            if (not selected_status or student['exam_status'] == selected_status)
            and (not selected_class_name or (student.get('class_name') or '').strip() == selected_class_name)
            and (not selected_major or (student.get('major') or '').strip() == selected_major)
            and (not selected_course or (student.get('teacher_name') or '').strip() == selected_course)
        ]
        class_options, major_options, course_options = _get_exam_setup_options()
        return render_template(
            'admin_exam_detail.html',
            exam=exam,
            batches=batches,
            selected_batch=selected_batch,
            selected_batch_id=selected_batch_id,
            students=visible_students,
            total_students=len(students),
            total_problems=total_problems,
            status_counts=status_counts,
            status_labels=status_labels,
            selected_status=selected_status,
            selected_class_name=selected_class_name,
            selected_major=selected_major,
            selected_course=selected_course,
            batch_class_options=batch_class_options,
            batch_major_options=batch_major_options,
            batch_course_options=batch_course_options,
            class_options=class_options,
            major_options=major_options,
            course_options=course_options,
        )


    @bp.route('/admin/exams/<int:exam_id>/batches/create', methods=['POST'])
    @login_required(db_check=True)
    def admin_create_exam_batch(exam_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
        exam = get_exam_by_id(exam_id)
        if not exam:
            flash('考试不存在或已被删除。', 'danger')
            return redirect(url_for('admin.admin_exams'))
        if exam.get('status') == 'archived':
            flash('归档考试只允许查看历史数据，不能再添加批次。', 'warning')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))

        batch_name = (request.form.get('batch_name') or '').strip()
        assign_type = (request.form.get('assign_type') or '').strip()
        assign_values = [value.strip() for value in request.form.getlist('assign_value') if value.strip()]
        upload_file = request.files.get('student_file')
        try:
            start_time = parse_exam_time((request.form.get('start_time') or '').strip())
            end_time = parse_exam_time((request.form.get('end_time') or '').strip())
        except ValueError as err:
            flash(str(err), 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))
        if not batch_name:
            flash('请填写批次名称。', 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))
        if start_time and end_time and end_time <= start_time:
            flash('结束时间必须晚于开始时间。', 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        paper_auto_enabled = False
        try:
            paper_auto_enabled = ensure_exam_paper_enabled(cursor, exam['paper_id'])
            missing_count = 0
            if assign_type == 'student':
                if not upload_file or not upload_file.filename or not allowed_excel_file(upload_file.filename):
                    raise ValueError('请上传包含学号或姓名的 .xlsx 文件。')
                assignments, missing_count = _resolve_uploaded_student_assignments(cursor, upload_file)
            else:
                if not assign_values:
                    raise ValueError('请选择批次的学生范围。')
                assignments = [(assign_type, value) for value in assign_values]
            batch_id, student_count = _create_exam_batch(
                cursor, exam_id, batch_name, start_time, end_time, assignments
            )
            conn.commit()
            clear_exam_metadata_cache()
            _finish_auto_enabled_paper(exam['paper_id'], paper_auto_enabled)
            message = f'批次创建成功，已分配 {student_count} 名学生。'
            if missing_count:
                message += f' 另有 {missing_count} 行未匹配到账户。'
            flash(message, 'success')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id, batch_id=batch_id))
        except (ValueError, mysql.connector.Error) as err:
            conn.rollback()
            flash(f'创建批次失败：{err}', 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))
        finally:
            cursor.close()
            conn.close()


    @bp.route('/admin/exams/<int:exam_id>/batches/<int:batch_id>/delete', methods=['POST'])
    @login_required(db_check=True)
    def admin_delete_exam_batch(exam_id, batch_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute(
                'SELECT * FROM exam_batches WHERE id = %s AND exam_id = %s FOR UPDATE',
                (batch_id, exam_id),
            )
            batch = cursor.fetchone()
            if not batch:
                raise ValueError('批次不存在。')
            if batch.get('status') == 'archived':
                raise ValueError('该批次已经归档。')
            cursor.execute('SELECT status FROM exams WHERE id = %s', (exam_id,))
            exam = cursor.fetchone()
            if not exam or exam.get('status') == 'archived':
                raise ValueError('归档考试不能再删除批次。')
            cursor.execute(
                """
                SELECT (
                    EXISTS(SELECT 1 FROM exam_user_questions WHERE batch_id = %s LIMIT 1)
                    OR EXISTS(SELECT 1 FROM user_responses WHERE batch_id = %s LIMIT 1)
                ) AS started
                """,
                (batch_id, batch_id),
            )
            if (cursor.fetchone() or {}).get('started'):
                cursor.execute("UPDATE exam_batches SET status = 'archived' WHERE id = %s", (batch_id,))
                message = f"批次《{batch['name']}》已有考试数据，已归档而不是删除。"
            else:
                cursor.execute('DELETE FROM exam_batches WHERE id = %s', (batch_id,))
                message = f"批次《{batch['name']}》已删除。"
            conn.commit()
            clear_exam_metadata_cache()
            flash(message, 'success')
        except (ValueError, mysql.connector.Error) as err:
            conn.rollback()
            flash(f'处理批次失败：{err}', 'danger')
        finally:
            cursor.close()
            conn.close()
        return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))


    @bp.route('/admin/exams/<int:exam_id>/batches/<int:batch_id>/extend', methods=['POST'])
    @login_required(db_check=True)
    def admin_extend_exam_batch(exam_id, batch_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
        try:
            new_end_time = parse_exam_time((request.form.get('end_time') or '').strip())
        except ValueError as err:
            flash(str(err), 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id, batch_id=batch_id))
        if not new_end_time:
            flash('请选择新的结束时间。', 'danger')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id, batch_id=batch_id))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute(
                'SELECT * FROM exam_batches WHERE id = %s AND exam_id = %s FOR UPDATE',
                (batch_id, exam_id),
            )
            batch = cursor.fetchone()
            if not batch:
                raise ValueError('批次不存在。')
            if batch.get('status') == 'archived':
                raise ValueError('归档批次不能再修改结束时间。')
            cursor.execute('SELECT status, paper_id FROM exams WHERE id = %s', (exam_id,))
            exam = cursor.fetchone()
            if not exam or exam.get('status') == 'archived':
                raise ValueError('归档考试不能再修改结束时间。')
            if batch.get('start_time') and new_end_time <= batch['start_time']:
                raise ValueError('结束时间必须晚于开始时间。')
            if batch.get('end_time') and new_end_time <= batch['end_time']:
                raise ValueError('只能延长结束时间，不能缩短。')
            paper_auto_enabled = ensure_exam_paper_enabled(cursor, exam['paper_id'])
            cursor.execute('UPDATE exam_batches SET end_time = %s WHERE id = %s', (new_end_time, batch_id))
            if batch.get('is_default'):
                cursor.execute('UPDATE exams SET end_time = %s WHERE id = %s', (new_end_time, exam_id))
            conn.commit()
            clear_exam_metadata_cache()
            _finish_auto_enabled_paper(exam['paper_id'], paper_auto_enabled)
            flash(f"批次《{batch['name']}》结束时间已延长。", 'success')
        except (ValueError, mysql.connector.Error) as err:
            conn.rollback()
            flash(f'延长考试时间失败：{err}', 'danger')
        finally:
            cursor.close()
            conn.close()
        return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id, batch_id=batch_id))


    @bp.route('/admin/exams/create', methods=['POST'])
    @login_required(db_check=True)
    def admin_create_exam():
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        name = (request.form.get('name') or '').strip()
        paper_id = request.form.get('paper_id', type=int)
        question_count = request.form.get('question_count', type=int)
        exam_type = (request.form.get('exam_type') or 'normal').strip()
        assign_type = (request.form.get('assign_type') or '').strip()
        assign_values = [
            value.strip()
            for value in request.form.getlist('assign_value')
            if value and value.strip()
        ]
        upload_file = request.files.get('student_file')
        batch_name = (request.form.get('batch_name') or '默认批次').strip()
        start_time_raw = (request.form.get('start_time') or '').strip()
        end_time_raw = (request.form.get('end_time') or '').strip()

        available_question_count = len(
            get_problem_templates_by_paper(paper_id, enabled_only=False)
        ) if paper_id else 0
        if not name or not paper_id or not question_count:
            flash('请填写考试名称、选择题库并设置抽题数量。', 'danger')
            return redirect(url_for('admin.admin_exams'))
        if question_count < 1 or question_count > available_question_count:
            flash(f'抽题数量必须在 1 到 {available_question_count} 之间。', 'danger')
            return redirect(url_for('admin.admin_exams'))
        if assign_type != 'student' and not assign_values:
            flash('请选择分配范围。', 'danger')
            return redirect(url_for('admin.admin_exams'))
        if assign_type == 'student':
            if not upload_file or upload_file.filename == '':
                flash('请上传包含学号或姓名的 Excel 文件。', 'danger')
                return redirect(url_for('admin.admin_exams'))
            if not allowed_excel_file(upload_file.filename):
                flash('文件格式错误，仅支持 .xlsx。', 'danger')
                return redirect(url_for('admin.admin_exams'))

        try:
            start_time = parse_exam_time(start_time_raw)
            end_time = parse_exam_time(end_time_raw)
        except ValueError as err:
            flash(str(err), 'danger')
            return redirect(url_for('admin.admin_exams'))
        if start_time and end_time and end_time <= start_time:
            flash('结束时间必须晚于开始时间。', 'danger')
            return redirect(url_for('admin.admin_exams'))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        paper_auto_enabled = False
        try:
            paper_auto_enabled = ensure_exam_paper_enabled(cursor, paper_id)
            cursor.execute(
                """
                INSERT INTO exams (name, paper_id, question_count, exam_type, start_time, end_time, status)
                VALUES (%s, %s, %s, %s, %s, %s, 'published')
                """,
                (name, paper_id, question_count, exam_type if exam_type in {'normal', 'retake'} else 'normal', start_time, end_time)
            )
            exam_id = cursor.lastrowid
            frozen_question_count = replace_exam_question_pool(cursor, exam_id, paper_id)
            if frozen_question_count < question_count:
                raise ValueError('当前题库可用题目不足，无法创建考试。')

            assignments = []
            missing_rows = []
            matched_ids = set()

            if assign_type == 'student':
                parsed_rows = _parse_exam_import_rows(upload_file)
                if not parsed_rows:
                    conn.rollback()
                    flash('Excel 中没有读取到可用的学号或姓名。', 'danger')
                    return redirect(url_for('admin.admin_exams'))

                seen_identifiers = set()
                for row in parsed_rows:
                    identifier = row['identifier']
                    student_id = row['student_id']
                    student_name = row['student_name']
                    lookup_key = f"{student_id}|{student_name}"
                    if lookup_key in seen_identifiers:
                        continue
                    seen_identifiers.add(lookup_key)

                    matched_user = None
                    match_reason = ''

                    if student_id:
                        cursor.execute(
                            "SELECT id, username, name FROM users WHERE username != 'admin' AND username = %s LIMIT 2",
                            (student_id,)
                        )
                        id_matches = cursor.fetchall()
                        if len(id_matches) == 1:
                            matched_user = id_matches[0]

                    if not matched_user and student_name:
                        cursor.execute(
                            "SELECT id, username, name FROM users WHERE username != 'admin' AND TRIM(name) = %s LIMIT 2",
                            (student_name,)
                        )
                        name_matches = cursor.fetchall()
                        if len(name_matches) == 1:
                            matched_user = name_matches[0]
                        elif len(name_matches) > 1:
                            match_reason = '姓名匹配到多个系统账户'

                    if matched_user:
                        matched_ids.add(int(matched_user['id']))
                    else:
                        if not match_reason:
                            if student_id and student_name:
                                match_reason = '学号和姓名均未找到匹配账户'
                            elif student_id:
                                match_reason = '学号未找到匹配账户'
                            else:
                                match_reason = '姓名未找到匹配账户'
                        missing_rows.append({
                            'row_no': row['row_no'],
                            'identifier': identifier,
                            'name': student_name,
                            'reason': match_reason,
                        })
            else:
                assignments = [(assign_type, assign_value) for assign_value in assign_values]

            if assign_type == 'student':
                if not matched_ids:
                    conn.rollback()
                    return _build_missing_accounts_response(missing_rows, name)
                assignments = [('student', str(user_id)) for user_id in sorted(matched_ids)]

            cursor.executemany(
                """
                INSERT IGNORE INTO exam_assignments (exam_id, assign_type, assign_value)
                VALUES (%s, %s, %s)
                """,
                [(exam_id, item_type, item_value) for item_type, item_value in assignments]
            )
            batch_id, batch_student_count = _create_exam_batch(
                cursor,
                exam_id,
                batch_name,
                start_time,
                end_time,
                assignments,
                is_default=True,
            )
            conn.commit()
            clear_exam_metadata_cache()
            _finish_auto_enabled_paper(paper_id, paper_auto_enabled)
            if assign_type == 'student' and missing_rows:
                return _build_missing_accounts_response(missing_rows, name)
            flash(f'考试创建成功，首个批次已分配 {batch_student_count} 名学生。', 'success')
        except ValueError as err:
            conn.rollback()
            flash(f'创建考试失败：{err}', 'danger')
        except mysql.connector.Error as err:
            conn.rollback()
            flash(f'创建考试失败：{err}', 'danger')
        finally:
            cursor.close()
            conn.close()

        return redirect(url_for('admin.admin_exams'))


    @bp.route('/admin/exams/<int:exam_id>/update', methods=['POST'])
    @login_required(db_check=True)
    def admin_update_exam(exam_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        existing_exam = get_exam_by_id(exam_id)
        if not existing_exam:
            flash('考试不存在或已被删除。', 'danger')
            return redirect(url_for('admin.admin_exams'))
        if existing_exam.get('status') == 'archived':
            flash('归档考试只允许查看历史数据，不能再编辑。', 'warning')
            return redirect(url_for('admin.admin_exam_detail', exam_id=exam_id))

        name = (request.form.get('name') or '').strip()
        paper_id = request.form.get('paper_id', type=int)
        question_count = request.form.get('question_count', type=int)
        exam_type = (request.form.get('exam_type') or 'normal').strip()
        assign_type = (request.form.get('assign_type') or '').strip()
        assign_values = [
            value.strip()
            for value in request.form.getlist('assign_value')
            if value and value.strip()
        ]
        upload_file = request.files.get('student_file')
        start_time_raw = (request.form.get('start_time') or '').strip()
        end_time_raw = (request.form.get('end_time') or '').strip()
        replace_student_assignments = False

        available_question_count = len(
            get_problem_templates_by_paper(paper_id, enabled_only=False)
        ) if paper_id else 0
        if not name or not paper_id or not question_count:
            flash('请填写考试名称、选择题库并设置抽题数量。', 'danger')
            return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
        if question_count < 1 or question_count > available_question_count:
            flash(f'抽题数量必须在 1 到 {available_question_count} 之间。', 'danger')
            return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
        if assign_type != 'student' and not assign_values:
            flash('请选择分配范围。', 'danger')
            return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
        if assign_type == 'student':
            if upload_file and upload_file.filename:
                if not allowed_excel_file(upload_file.filename):
                    flash('文件格式错误，仅支持 .xlsx。', 'danger')
                    return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
                replace_student_assignments = True
            else:
                cursor = None
                conn = get_db_connection()
                cursor = conn.cursor(dictionary=True)
                try:
                    cursor.execute(
                        "SELECT COUNT(*) AS total FROM exam_assignments WHERE exam_id = %s AND assign_type = 'student'",
                        (exam_id,)
                    )
                    has_existing_assignments = (cursor.fetchone() or {}).get('total', 0) > 0
                finally:
                    cursor.close()
                    conn.close()
                if not has_existing_assignments:
                    flash('当前考试没有可保留的学生名单，请上传 Excel 后再保存。', 'danger')
                    return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
        elif upload_file and upload_file.filename:
            if not allowed_excel_file(upload_file.filename):
                flash('文件格式错误，仅支持 .xlsx。', 'danger')
                return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))

        try:
            start_time = parse_exam_time(start_time_raw)
            end_time = parse_exam_time(end_time_raw)
        except ValueError as err:
            flash(str(err), 'danger')
            return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
        if start_time and end_time and end_time <= start_time:
            flash('结束时间必须晚于开始时间。', 'danger')
            return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        paper_auto_enabled = False
        try:
            paper_auto_enabled = ensure_exam_paper_enabled(cursor, paper_id)
            cursor.execute(
                "SELECT * FROM exam_batches WHERE exam_id = %s AND is_default = TRUE LIMIT 1 FOR UPDATE",
                (exam_id,),
            )
            default_batch = cursor.fetchone()
            if not default_batch:
                raise ValueError('默认批次不存在，请先重新部署以完成数据迁移。')
            cursor.execute(
                "SELECT assign_type, assign_value FROM exam_batch_assignments WHERE batch_id = %s",
                (default_batch['id'],),
            )
            current_batch_assignments = cursor.fetchall()
            current_assignment_type = current_batch_assignments[0]['assign_type'] if current_batch_assignments else ''
            current_assignment_values = {
                row['assign_value'] for row in current_batch_assignments
                if row['assign_type'] == current_assignment_type
            }
            if assign_type == 'student' and not replace_student_assignments:
                assignment_changed = False
            else:
                assignment_changed = (
                    assign_type != current_assignment_type or set(assign_values) != current_assignment_values
                )
            cursor.execute(
                "SELECT EXISTS(SELECT 1 FROM exam_user_questions WHERE batch_id = %s LIMIT 1) AS started",
                (default_batch['id'],),
            )
            default_batch_started = bool((cursor.fetchone() or {}).get('started'))
            if default_batch_started:
                start_changed = start_time != default_batch.get('start_time')
                old_end_time = default_batch.get('end_time')
                end_shortened = bool(old_end_time and end_time and end_time < old_end_time)
                end_removed_or_added = (old_end_time is None) != (end_time is None)
                if start_changed or assignment_changed or end_shortened or end_removed_or_added:
                    conn.rollback()
                    flash('默认批次已有学生进入，只能保持名单和开始时间不变，并向后延长结束时间。', 'danger')
                    return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))

            old_question_count = existing_exam.get('question_count')
            if old_question_count is None and int(existing_exam['paper_id']) == paper_id:
                old_paper_count = len(
                    get_problem_templates_by_paper(paper_id, enabled_only=False)
                )
                stored_question_count = None if question_count == old_paper_count else question_count
            else:
                stored_question_count = question_count
            question_config_changed = (
                int(existing_exam['paper_id']) != paper_id or
                old_question_count != stored_question_count
            )
            if question_config_changed:
                cursor.execute(
                    """
                    SELECT (
                        EXISTS(SELECT 1 FROM exam_user_questions WHERE exam_id = %s LIMIT 1)
                        OR EXISTS(SELECT 1 FROM user_responses WHERE exam_id = %s LIMIT 1)
                    ) AS started
                    """,
                    (exam_id, exam_id),
                )
                if (cursor.fetchone() or {}).get('started'):
                    conn.rollback()
                    flash('已有学生进入过答题页面，不能再修改题库或抽题数量。', 'danger')
                    return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
                cursor.execute('DELETE FROM exam_user_questions WHERE exam_id = %s', (exam_id,))
                frozen_question_count = replace_exam_question_pool(cursor, exam_id, paper_id)
                if frozen_question_count < question_count:
                    raise ValueError('当前题库可用题目不足，无法保存考试。')

            cursor.execute(
                """
                UPDATE exams
                SET name = %s,
                    paper_id = %s,
                    question_count = %s,
                    exam_type = %s,
                    start_time = %s,
                    end_time = %s,
                    status = 'published'
                WHERE id = %s
                """,
                (name, paper_id, stored_question_count, exam_type if exam_type in {'normal', 'retake'} else 'normal', start_time, end_time, exam_id)
            )

            assignments = []
            missing_rows = []
            matched_ids = set()

            if assign_type == 'student':
                if replace_student_assignments:
                    parsed_rows = _parse_exam_import_rows(upload_file)
                    if not parsed_rows:
                        conn.rollback()
                        flash('Excel 中没有读取到可用的学号或姓名。', 'danger')
                        return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))

                    seen_identifiers = set()
                    for row in parsed_rows:
                        identifier = row['identifier']
                        student_id = row['student_id']
                        student_name = row['student_name']
                        lookup_key = f"{student_id}|{student_name}"
                        if lookup_key in seen_identifiers:
                            continue
                        seen_identifiers.add(lookup_key)

                        matched_user = None
                        match_reason = ''

                        if student_id:
                            cursor.execute(
                                "SELECT id, username, name FROM users WHERE username != 'admin' AND username = %s LIMIT 2",
                                (student_id,)
                            )
                            id_matches = cursor.fetchall()
                            if len(id_matches) == 1:
                                matched_user = id_matches[0]

                        if not matched_user and student_name:
                            cursor.execute(
                                "SELECT id, username, name FROM users WHERE username != 'admin' AND TRIM(name) = %s LIMIT 2",
                                (student_name,)
                            )
                            name_matches = cursor.fetchall()
                            if len(name_matches) == 1:
                                matched_user = name_matches[0]
                            elif len(name_matches) > 1:
                                match_reason = '姓名匹配到多个系统账户'

                        if matched_user:
                            matched_ids.add(int(matched_user['id']))
                        else:
                            if not match_reason:
                                if student_id and student_name:
                                    match_reason = '学号和姓名均未找到匹配账户'
                                elif student_id:
                                    match_reason = '学号未找到匹配账户'
                                else:
                                    match_reason = '姓名未找到匹配账户'
                            missing_rows.append({
                                'row_no': row['row_no'],
                                'identifier': identifier,
                                'name': student_name,
                                'reason': match_reason,
                            })
                    cursor.execute('DELETE FROM exam_assignments WHERE exam_id = %s', (exam_id,))
                else:
                    cursor.execute(
                        'UPDATE exam_batches SET start_time = %s, end_time = %s WHERE id = %s',
                        (start_time, end_time, default_batch['id']),
                    )
                    conn.commit()
                    clear_exam_metadata_cache()
                    _finish_auto_enabled_paper(paper_id, paper_auto_enabled)
                    flash('考试更新成功，已保留原有学生名单。', 'success')
                    return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))
            else:
                cursor.execute('DELETE FROM exam_assignments WHERE exam_id = %s', (exam_id,))
                assignments = [(assign_type, assign_value) for assign_value in assign_values]

            if assign_type == 'student':
                if not matched_ids:
                    conn.rollback()
                    return _build_missing_accounts_response(missing_rows, name)
                assignments = [('student', str(user_id)) for user_id in sorted(matched_ids)]

            cursor.executemany(
                """
                INSERT IGNORE INTO exam_assignments (exam_id, assign_type, assign_value)
                VALUES (%s, %s, %s)
                """,
                [(exam_id, item_type, item_value) for item_type, item_value in assignments]
            )
            cursor.execute(
                'UPDATE exam_batches SET start_time = %s, end_time = %s WHERE id = %s',
                (start_time, end_time, default_batch['id']),
            )
            if assignment_changed:
                _replace_exam_batch_assignments(
                    cursor, exam_id, default_batch['id'], assignments
                )
            conn.commit()
            clear_exam_metadata_cache()
            _finish_auto_enabled_paper(paper_id, paper_auto_enabled)
            if assign_type == 'student' and missing_rows:
                return _build_missing_accounts_response(missing_rows, name)
            flash(f'考试更新成功，已设置 {len(assignments)} 条分配规则。', 'success')
        except (ValueError, mysql.connector.Error) as err:
            conn.rollback()
            flash(f'更新考试失败：{err}', 'danger')
        finally:
            cursor.close()
            conn.close()

        return redirect(url_for('admin.admin_exams', edit_exam_id=exam_id))


    @bp.route('/admin/exams/retake_candidates')
    @login_required
    def admin_retake_candidates():
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        source_exam_id = request.args.get('source_exam_id', type=int)
        status_filter = (request.args.get('status') or 'incomplete').strip()
        candidates = []
        source_exam = get_exam_by_id(source_exam_id) if source_exam_id else None
        if source_exam:
            conn = get_db_connection()
            cursor = conn.cursor(dictionary=True)
            try:
                total_problems = (
                    source_exam['question_count']
                    if source_exam.get('question_count') is not None
                    else get_exam_question_pool_count(source_exam['id'], source_exam['paper_id'])
                )
                cursor.execute(
                    """
                    SELECT u.id, u.username, u.name, u.class_name, u.major, u.teacher_name,
                           COALESCE(c.completed_count, 0) AS completed_count
                    FROM users u
                    LEFT JOIN (
                        SELECT user_id, COUNT(DISTINCT template_id) AS completed_count
                        FROM (
                            SELECT user_id, template_id, attempt_count
                            FROM user_responses
                            WHERE exam_id = %s
                            GROUP BY user_id, template_id, attempt_count
                            HAVING SUM(CASE WHEN is_correct THEN 1 ELSE 0 END) = COUNT(*)
                        ) good
                        GROUP BY user_id
                    ) c ON c.user_id = u.id
                    WHERE u.username != 'admin'
                      AND EXISTS (
                        SELECT 1
                        FROM exam_batch_students s
                        WHERE s.exam_id = %s AND s.user_id = u.id
                      )
                    HAVING completed_count < %s
                    ORDER BY u.class_name, u.username
                    LIMIT 1000
                    """,
                    (source_exam_id, source_exam_id, total_problems)
                )
                candidates = cursor.fetchall()
            finally:
                cursor.close()
                conn.close()
        elif source_exam_id:
            flash('原考试不存在。', 'danger')

        class_options, major_options, course_options = _get_exam_setup_options()
        return render_template(
            'admin_exams.html',
            exams=get_all_exams(),
            exam_papers=get_exam_papers(include_disabled=True),
            class_options=class_options,
            major_options=major_options,
            course_options=course_options,
            retake_candidates=candidates,
            source_exam=source_exam,
            selected_source_exam_id=source_exam_id,
            status_filter=status_filter,
            format_datetime_local=format_datetime_local
        )
    
    
    @bp.route('/admin/import/students', methods=['POST'])
    @login_required(db_check=True)
    def admin_import_students():
        """管理员批量导入学生（xlsx：学号、姓名、专业、班级、课程号）。"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
    
        upload_file = request.files.get('students_file')
        if not upload_file or upload_file.filename == '':
            flash('请先选择要导入的 xlsx 文件', 'danger')
            return redirect(url_for('admin_dashboard'))
    
        if not allowed_excel_file(upload_file.filename):
            flash('文件格式错误，仅支持 .xlsx', 'danger')
            return redirect(url_for('admin_dashboard'))
    
        default_teacher_name = (request.form.get('teacher_name') or '').strip()
    
        try:
            workbook = load_workbook(upload_file, read_only=True, data_only=True)
            sheet = workbook.active
    
            raw_rows = []
            for row in sheet.iter_rows(values_only=True):
                student_id = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ''
                student_name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
                major = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''
                class_name = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ''
                teacher_name = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ''
                if not student_id and not student_name and not major and not class_name and not teacher_name:
                    continue
                raw_rows.append((student_id, student_name, major, class_name, teacher_name))
        except Exception as e:
            flash(f'读取 xlsx 失败：{e}', 'danger')
            return redirect(url_for('admin_dashboard'))
    
        if not raw_rows:
            flash('未读取到有效数据，请检查文件内容', 'warning')
            return redirect(url_for('admin_dashboard'))
    
        normalized_rows = raw_rows
        first_id, first_name, first_major, first_class, first_teacher = raw_rows[0]
        if first_id.lower() in {'学号', 'student_id', 'studentid', 'id', '账号', '用户名'}:
            if (first_name.lower() in {'姓名', 'name', '学生姓名'} or not first_name) and \
                    (first_major.lower() in {'专业', 'major', 'major_name'} or not first_major) and \
                    (first_class.lower() in {'班级', 'class', 'class_name'} or not first_class) and \
                    (first_teacher.lower() in {'老师', '教师', '任课老师', '课程号', 'course_number', 'teacher', 'teacher_name'} or not first_teacher):
                normalized_rows = raw_rows[1:]
    
        valid_rows = []
        skipped_rows = 0
        for student_id, student_name, major, class_name, teacher_name in normalized_rows:
            if not student_id:
                skipped_rows += 1
                continue
            teacher_name = teacher_name or default_teacher_name
            valid_rows.append((student_id, student_name or None, major or None, class_name or None, teacher_name or None))
    
        if not valid_rows:
            flash('未找到可导入的学号数据', 'warning')
            return redirect(url_for('admin_dashboard'))
    
        conn = get_db_connection()
        if not conn:
            flash('数据库连接失败', 'danger')
            return redirect(url_for('admin_dashboard'))
    
        inserted_count = 0
        updated_count = 0
        cursor = None
        try:
            cursor = conn.cursor()
            for student_id, student_name, major, class_name, teacher_name in valid_rows:
                initial_password = build_initial_password(student_id)
                initial_password_hash = generate_password_hash(initial_password)
                cursor.execute("SELECT id FROM users WHERE username = %s", (student_id,))
                existing = cursor.fetchone()
                if existing:
                    cursor.execute(
                        """
                        UPDATE users
                        SET name = %s,
                            major = %s,
                            class_name = %s,
                            teacher_name = %s,
                            password = %s,
                            password_changed = FALSE
                        WHERE username = %s
                        """,
                        (student_name, major, class_name, teacher_name, initial_password_hash, student_id)
                    )
                    updated_count += 1
                else:
                    cursor.execute(
                        """
                        INSERT INTO users (username, name, major, class_name, teacher_name, password, password_changed, avatar_filename)
                        VALUES (%s, %s, %s, %s, %s, %s, FALSE, %s)
                        """,
                        (student_id, student_name, major, class_name, teacher_name, initial_password_hash, DEFAULT_AVATAR)
                    )
                    inserted_count += 1
    
            conn.commit()
            flash(
                f'导入完成：新增 {inserted_count} 人，更新 {updated_count} 人，跳过 {skipped_rows} 行。'
                f'初始密码已按“@ncst+学号后四位”自动设置。',
                'success'
            )
        except Exception as e:
            conn.rollback()
            flash(f'导入失败：{e}', 'danger')
        finally:
            if cursor:
                cursor.close()
            conn.close()
    
        return redirect(url_for('admin_dashboard'))
    
    
    @bp.route('/admin/export/all-student-analytics')
    @login_required
    def admin_export_all_student_analytics():
        """导出所有学生完整答题分析 Excel。"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        selected_exam_id = request.args.get('exam_id', type=int)
        selected_exam = get_exam_by_id(selected_exam_id) if selected_exam_id else None
        if selected_exam_id and not selected_exam:
            flash('考试不存在，无法导出答题数据。', 'danger')
            return redirect(url_for('admin.admin_exams'))

        if selected_exam:
            selected_paper_id = selected_exam['paper_id']
            display_mapping = get_problem_display_info(selected_paper_id, enabled_only=False)
            response_filter = " AND ur.exam_id = %s"
            response_params = [selected_exam_id]
        else:
            selected_paper_id = resolve_selected_exam_paper_id(
                request.args.get('paper_id', type=int),
                include_disabled_for_admin=True
            )
            display_mapping = get_problem_display_info(selected_paper_id, enabled_only=True)
            response_filter, response_params = build_enabled_paper_filter('ur', selected_paper_id)

        problem_items = sorted(display_mapping.values(), key=lambda item: item['display_number'])
        total_problems = len(problem_items)
        status = (request.args.get('status') or '').strip()
        keyword = (request.args.get('keyword') or '').strip()
        class_name = (request.args.get('class_name') or '').strip()
        major = (request.args.get('major') or '').strip()
        teacher_name = (request.args.get('teacher_name') or '').strip()
        status_filter_requested = not selected_exam and status in {'completed', 'incomplete'}
        has_student_filter = bool(keyword or class_name or major or teacher_name)
        if status_filter_requested and not has_student_filter:
            flash('请先筛选学生后再导出。', 'warning')
            return redirect(url_for('admin_students_by_status', status=status, paper_id=selected_paper_id))

        filtered_export = status_filter_requested
        selected_students = None
        selected_user_ids = []

        if filtered_export:
            selected_students = get_students_by_completion(
                completed=(status == 'completed'),
                paper_id=selected_paper_id,
                filters={
                    'keyword': keyword,
                    'class_name': class_name,
                    'major': major,
                    'teacher_name': teacher_name,
                }
            )
            selected_user_ids = [student['id'] for student in selected_students]

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            if selected_students is not None:
                students = selected_students
            elif selected_exam:
                cursor.execute("""
                    SELECT DISTINCT u.id, u.username, u.name, u.major, u.class_name, u.teacher_name,
                           u.exam_start_time, u.exam_end_time, u.completed_at, u.created_at
                    FROM users u
                    WHERE u.username != 'admin'
                      AND EXISTS (
                        SELECT 1
                        FROM exam_batch_students s
                        WHERE s.exam_id = %s AND s.user_id = u.id
                      )
                    ORDER BY COALESCE(NULLIF(TRIM(u.teacher_name), ''), '未分配') ASC,
                             COALESCE(NULLIF(TRIM(u.class_name), ''), '未分班') ASC,
                             u.username ASC
                """, (selected_exam_id,))
                students = cursor.fetchall()
            else:
                cursor.execute("""
                    SELECT id, username, name, major, class_name, teacher_name, exam_start_time, exam_end_time,
                           completed_at, created_at
                    FROM users
                    WHERE username != 'admin'
                    ORDER BY COALESCE(NULLIF(TRIM(teacher_name), ''), '未分配') ASC,
                             COALESCE(NULLIF(TRIM(class_name), ''), '未分班') ASC,
                             username ASC
                """)
                students = cursor.fetchall()

            user_filter_sql = ""
            user_filter_params = []
            if selected_exam:
                selected_user_ids = [student['id'] for student in students]

            if selected_exam or filtered_export:
                if selected_user_ids:
                    user_placeholders = ', '.join(['%s'] * len(selected_user_ids))
                    user_filter_sql = f" AND ur.user_id IN ({user_placeholders})"
                    user_filter_params = selected_user_ids
                else:
                    user_filter_sql = " AND 1 = 0"

            cursor.execute(f"""
                WITH attempt_summary AS (
                    SELECT
                        ur.user_id,
                        ur.template_id,
                        ur.attempt_count,
                        COUNT(*) AS answer_rows,
                        SUM(CASE WHEN ur.is_correct THEN 1 ELSE 0 END) AS correct_rows,
                        CASE WHEN SUM(CASE WHEN ur.is_correct THEN 1 ELSE 0 END) = COUNT(*) THEN 1 ELSE 0 END AS is_fully_correct,
                        MAX(ur.time_taken) AS time_taken,
                        MAX(COALESCE(ur.switch_count, 0)) AS switch_count,
                        MAX(ur.response_time) AS last_response_time
                    FROM user_responses ur
                    WHERE 1 = 1 {response_filter} {user_filter_sql}
                    GROUP BY ur.user_id, ur.template_id, ur.attempt_count
                )
                SELECT
                    user_id,
                    template_id,
                    COUNT(*) AS total_attempts,
                    SUM(is_fully_correct) AS correct_attempts,
                    MAX(is_fully_correct) AS is_completed,
                    MIN(CASE WHEN is_fully_correct = 1 THEN attempt_count ELSE NULL END) AS attempts_to_correct,
                    SUM(time_taken) AS total_time_spent,
                    SUM(switch_count) AS total_switch_count,
                    MAX(last_response_time) AS last_response_time,
                    SUM(CASE WHEN is_fully_correct = 0 THEN 1 ELSE 0 END) AS wrong_attempts
                FROM attempt_summary
                GROUP BY user_id, template_id
            """, response_params + user_filter_params)
            per_problem_rows = cursor.fetchall()

            per_student_problem = {
                (row['user_id'], row['template_id']): row
                for row in per_problem_rows
            }

            cursor.execute(f"""
                WITH attempt_summary AS (
                    SELECT
                        ur.user_id,
                        ur.template_id,
                        ur.attempt_count,
                        COUNT(*) AS answer_rows,
                        SUM(CASE WHEN ur.is_correct THEN 1 ELSE 0 END) AS correct_rows,
                        CASE WHEN SUM(CASE WHEN ur.is_correct THEN 1 ELSE 0 END) = COUNT(*) THEN 1 ELSE 0 END AS is_fully_correct,
                        MAX(ur.time_taken) AS time_taken,
                        MAX(COALESCE(ur.switch_count, 0)) AS switch_count,
                        MAX(ur.response_time) AS last_response_time
                    FROM user_responses ur
                    WHERE 1 = 1 {response_filter} {user_filter_sql}
                    GROUP BY ur.user_id, ur.template_id, ur.attempt_count
                )
                SELECT
                    user_id,
                    COUNT(*) AS total_attempts,
                    SUM(is_fully_correct) AS correct_attempts,
                    COUNT(DISTINCT CASE WHEN is_fully_correct = 1 THEN template_id ELSE NULL END) AS completed_problem_count,
                    SUM(time_taken) AS total_time_spent,
                    SUM(switch_count) AS total_switch_count,
                    MAX(last_response_time) AS last_response_time
                FROM attempt_summary
                GROUP BY user_id
            """, response_params + user_filter_params)
            summary_rows = {row['user_id']: row for row in cursor.fetchall()}

            cursor.execute(f"""
                SELECT
                    ur.user_id,
                    ur.template_id,
                    ur.attempt_count,
                    ur.answer_index,
                    ur.user_answer,
                    ur.correct_answer,
                    ur.is_correct,
                    ur.error_type,
                    ur.time_taken,
                    COALESCE(ur.switch_count, 0) AS switch_count,
                    ur.response_time
                FROM user_responses ur
                JOIN users u ON u.id = ur.user_id
                WHERE u.username != 'admin' {response_filter} {user_filter_sql}
                ORDER BY COALESCE(NULLIF(TRIM(u.teacher_name), ''), '未分配') ASC,
                         COALESCE(NULLIF(TRIM(u.class_name), ''), '未分班') ASC,
                         u.username ASC,
                         ur.template_id ASC,
                         ur.attempt_count ASC,
                         ur.answer_index ASC
            """, response_params + user_filter_params)
            raw_rows = cursor.fetchall()
        finally:
            cursor.close()
            conn.close()

        def as_float(value):
            if value is None:
                return 0
            return float(value)

        def fmt_dt(value):
            return value.strftime('%Y-%m-%d %H:%M:%S') if value else ''

        wb = Workbook()
        summary_ws = wb.active
        summary_ws.title = '学生汇总宽表'
        detail_ws = wb.create_sheet('逐题明细长表')
        raw_ws = wb.create_sheet('原始答题记录')

        header_fill = PatternFill('solid', fgColor='1F4E78')
        header_font = Font(color='FFFFFF', bold=True)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)

        base_headers = [
            '课程号', '班级', '专业', '学号', '姓名',
            '开考时间', '结束时间', '完成时间',
            '答对题数', '总题数', '完成率(%)',
            '答题次数', '答对次数', '总体正确率(%)',
            '累计用时(秒)', '切屏次数'
        ]
        problem_headers = []
        for item in problem_items:
            prefix = f"第{item['display_number']}题"
            problem_headers.extend([
                f'{prefix}-答题次数',
                f'{prefix}-是否完成',
                f'{prefix}-答对次数',
                f'{prefix}-错误次数',
                f'{prefix}-首次答对第几次',
                f'{prefix}-累计时长(秒)',
                f'{prefix}-切屏次数',
                f'{prefix}-最后答题时间',
            ])
        summary_ws.append(base_headers + problem_headers)

        for student in students:
            summary = summary_rows.get(student['id'], {})
            completed_count = int(summary.get('completed_problem_count') or 0)
            total_attempts = int(summary.get('total_attempts') or 0)
            correct_attempts = int(summary.get('correct_attempts') or 0)
            completion_rate = round(completed_count / total_problems * 100, 1) if total_problems else 0
            overall_rate = round(correct_attempts / total_attempts * 100, 1) if total_attempts else 0
            start_time_value = selected_exam.get('start_time') if selected_exam else student.get('exam_start_time')
            end_time_value = selected_exam.get('end_time') if selected_exam else student.get('exam_end_time')
            completed_at_value = (
                summary.get('last_response_time')
                if selected_exam and total_problems and completed_count >= total_problems
                else student.get('completed_at')
            )
            row = [
                student.get('teacher_name') or '未分配',
                student.get('class_name') or '未分班',
                student.get('major') or '未设置专业',
                student.get('username') or '',
                student.get('name') or '',
                fmt_dt(start_time_value),
                fmt_dt(end_time_value),
                fmt_dt(completed_at_value),
                completed_count,
                total_problems,
                completion_rate,
                total_attempts,
                correct_attempts,
                overall_rate,
                round(as_float(summary.get('total_time_spent')), 1),
                int(summary.get('total_switch_count') or 0),
            ]
            for item in problem_items:
                stat = per_student_problem.get((student['id'], item['actual_id']), {})
                attempts = int(stat.get('total_attempts') or 0)
                completed_flag = bool(stat.get('is_completed'))
                row.extend([
                    attempts,
                    '是' if completed_flag else '否',
                    int(stat.get('correct_attempts') or 0),
                    int(stat.get('wrong_attempts') or 0),
                    stat.get('attempts_to_correct') or '',
                    round(as_float(stat.get('total_time_spent')), 1),
                    int(stat.get('total_switch_count') or 0),
                    fmt_dt(stat.get('last_response_time')),
                ])
            summary_ws.append(row)

        detail_ws.append([
            '课程号', '班级', '专业', '学号', '姓名',
            '题号', '题目名称', '答题次数', '是否完成', '答对次数', '错误次数',
            '首次答对第几次', '累计时长(秒)', '切屏次数', '最后答题时间'
        ])
        problem_name_by_id = {item['actual_id']: item['template_name'] for item in problem_items}
        problem_no_by_id = {item['actual_id']: item['display_number'] for item in problem_items}
        student_username_by_id = {student['id']: student.get('username') or '' for student in students}
        for student in students:
            for item in problem_items:
                stat = per_student_problem.get((student['id'], item['actual_id']), {})
                detail_ws.append([
                    student.get('teacher_name') or '未分配',
                    student.get('class_name') or '未分班',
                    student.get('major') or '未设置专业',
                    student.get('username') or '',
                    student.get('name') or '',
                    item['display_number'],
                    item['template_name'],
                    int(stat.get('total_attempts') or 0),
                    '是' if stat.get('is_completed') else '否',
                    int(stat.get('correct_attempts') or 0),
                    int(stat.get('wrong_attempts') or 0),
                    stat.get('attempts_to_correct') or '',
                    round(as_float(stat.get('total_time_spent')), 1),
                    int(stat.get('total_switch_count') or 0),
                    fmt_dt(stat.get('last_response_time')),
                ])

        raw_ws.append([
            '学号', '题号', '题目名称', '尝试次数', '答案序号', '学生答案', '正确答案',
            '是否正确', '错误类型', '本次用时(秒)', '本次切屏次数', '提交时间'
        ])
        for row in raw_rows:
            raw_ws.append([
                student_username_by_id.get(row['user_id'], ''),
                problem_no_by_id.get(row['template_id'], row['template_id']),
                problem_name_by_id.get(row['template_id'], ''),
                row.get('attempt_count') or 0,
                (row.get('answer_index') or 0) + 1,
                row.get('user_answer'),
                row.get('correct_answer'),
                '是' if row.get('is_correct') else '否',
                row.get('error_type') or '',
                round(as_float(row.get('time_taken')), 1),
                int(row.get('switch_count') or 0),
                fmt_dt(row.get('response_time')),
            ])

        for ws in wb.worksheets:
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center
            for column_cells in ws.columns:
                column_letter = get_column_letter(column_cells[0].column)
                max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
                ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 24)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        if selected_exam:
            safe_exam_name = re.sub(r'[^A-Za-z0-9_-]+', '_', selected_exam.get('name') or '').strip('_')
            filename_prefix = f"exam_{selected_exam_id}_{safe_exam_name or 'student_answer_analytics'}"
        else:
            filename_prefix = f"{status}_filtered_student_answer_analytics" if filtered_export else "all_student_answer_analytics"
        filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response


    @bp.route('/admin/exams/<int:exam_id>/delete', methods=['POST'])
    @login_required(db_check=True)
    def admin_delete_exam(exam_id):
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute('SELECT name FROM exams WHERE id = %s FOR UPDATE', (exam_id,))
            exam = cursor.fetchone()
            if not exam:
                flash('考试不存在或已被删除。', 'warning')
                return redirect(url_for('admin.admin_exams'))

            cursor.execute(
                """
                SELECT (
                    EXISTS(SELECT 1 FROM exam_user_questions WHERE exam_id = %s LIMIT 1)
                    OR EXISTS(SELECT 1 FROM user_responses WHERE exam_id = %s LIMIT 1)
                ) AS started
                """,
                (exam_id, exam_id),
            )
            if (cursor.fetchone() or {}).get('started'):
                cursor.execute("UPDATE exams SET status = 'archived' WHERE id = %s", (exam_id,))
                cursor.execute("UPDATE exam_batches SET status = 'archived' WHERE exam_id = %s", (exam_id,))
                message = f"考试《{exam['name']}》已有考试数据，已归档并保留全部成绩。"
            else:
                cursor.execute('DELETE FROM exam_assignments WHERE exam_id = %s', (exam_id,))
                cursor.execute('DELETE FROM exam_question_pool WHERE exam_id = %s', (exam_id,))
                cursor.execute('DELETE FROM exams WHERE id = %s', (exam_id,))
                message = f"考试《{exam['name']}》已删除。"
            conn.commit()
            clear_exam_metadata_cache()
            flash(message, 'success')
        except mysql.connector.Error as err:
            conn.rollback()
            flash(f'删除考试失败：{err}', 'danger')
        finally:
            cursor.close()
            conn.close()

        return redirect(url_for('admin.admin_exams'))
    

    @bp.route('/admin/export/best-exam-scores')
    @login_required
    def admin_export_best_exam_scores():
        """导出当前开启题库中的学生最高答对题数。"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))

        enabled_papers = get_enabled_exam_papers()
        if not enabled_papers:
            flash('当前没有开启的题库，无法导出最高成绩。', 'warning')
            return redirect(url_for('admin_dashboard'))

        paper_ids = [paper['id'] for paper in enabled_papers]
        paper_by_id = {paper['id']: paper for paper in enabled_papers}
        placeholders = ', '.join(['%s'] * len(paper_ids))

        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            cursor.execute("""
                SELECT id, username, name, major, class_name, teacher_name
                FROM users
                WHERE username != 'admin'
                ORDER BY COALESCE(NULLIF(TRIM(teacher_name), ''), '未分配') ASC,
                         COALESCE(NULLIF(TRIM(class_name), ''), '未分班') ASC,
                         username ASC
            """)
            students = cursor.fetchall()

            cursor.execute(f"""
                SELECT paper_id, COUNT(*) AS total_problems
                FROM problem_templates
                WHERE paper_id IN ({placeholders})
                  AND COALESCE(status, 'active') = 'active'
                GROUP BY paper_id
            """, paper_ids)
            total_problem_by_paper = {
                row['paper_id']: int(row.get('total_problems') or 0)
                for row in cursor.fetchall()
            }

            cursor.execute(f"""
                WITH attempt_summary AS (
                    SELECT
                        ur.user_id,
                        COALESCE(ur.paper_id, pt.paper_id) AS response_paper_id,
                        ur.template_id,
                        ur.attempt_count,
                        CASE WHEN SUM(CASE WHEN ur.is_correct THEN 1 ELSE 0 END) = COUNT(*) THEN 1 ELSE 0 END AS is_fully_correct
                    FROM user_responses ur
                    JOIN problem_templates pt ON pt.id = ur.template_id
                    WHERE COALESCE(ur.paper_id, pt.paper_id) IN ({placeholders})
                    GROUP BY ur.user_id, response_paper_id, ur.template_id, ur.attempt_count
                ),
                template_summary AS (
                    SELECT
                        user_id,
                        response_paper_id,
                        template_id,
                        MAX(is_fully_correct) AS template_correct
                    FROM attempt_summary
                    GROUP BY user_id, response_paper_id, template_id
                )
                SELECT
                    user_id,
                    response_paper_id AS paper_id,
                    SUM(template_correct) AS correct_count
                FROM template_summary
                GROUP BY user_id, response_paper_id
            """, paper_ids)
            score_rows = cursor.fetchall()
        finally:
            cursor.close()
            conn.close()

        score_by_user_paper = {
            (row['user_id'], row['paper_id']): int(row.get('correct_count') or 0)
            for row in score_rows
        }

        wb = Workbook()
        summary_ws = wb.active
        summary_ws.title = '最高成绩汇总'

        header_fill = PatternFill('solid', fgColor='1F4E78')
        header_font = Font(color='FFFFFF', bold=True)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)

        paper_headers = [paper['name'] for paper in enabled_papers]
        summary_ws.append(
            ['课程号', '班级', '专业', '学号', '姓名']
            + paper_headers
            + ['最高得分', '最高得分题库']
        )

        for student in students:
            paper_scores = []
            best_score = 0
            best_paper_name = ''
            for paper in enabled_papers:
                score = score_by_user_paper.get((student['id'], paper['id']), 0)
                paper_scores.append(score)
                if score > best_score:
                    best_score = score
                    best_paper_name = paper['name']

            summary_ws.append([
                student.get('teacher_name') or '未分配',
                student.get('class_name') or '未分班',
                student.get('major') or '未设置专业',
                student.get('username') or '',
                student.get('name') or '',
            ] + paper_scores + [best_score, best_paper_name])

        info_ws = wb.create_sheet('题库信息')
        info_ws.append(['题库名称', '题目总数'])
        for paper_id in paper_ids:
            info_ws.append([
                paper_by_id[paper_id]['name'],
                total_problem_by_paper.get(paper_id, 0)
            ])

        for ws in wb.worksheets:
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center
            for column_cells in ws.columns:
                column_letter = get_column_letter(column_cells[0].column)
                max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
                ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 24)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename_prefix = "best_scores_enabled_papers"
        filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        return response
    
    

    @bp.route('/admin/image_manager')
    @login_required
    def admin_image_manager():
        """图片管理页面"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
    
        import os
        from datetime import datetime
    
        images = []
        images_path = os.path.join(app.root_path, 'static', 'images')
    
        if os.path.exists(images_path):
            for filename in os.listdir(images_path):
                if allowed_file(filename):
                    file_path = os.path.join(images_path, filename)
                    file_stat = os.stat(file_path)
                    images.append({
                        'filename': filename,
                        'size': round(file_stat.st_size / 1024, 1),  # KB
                        'modified': datetime.fromtimestamp(file_stat.st_mtime).strftime('%Y-%m-%d %H:%M')
                    })
    
        return render_template('admin_image_manager.html', images=images)
    
    
    @bp.route('/admin/delete_image/<filename>')
    @login_required(db_check=True)
    def admin_delete_image(filename):
        """删除图片"""
        if session.get('username') != 'admin':
            flash('权限不足', 'danger')
            return redirect(url_for('dashboard'))
    
        import os
        from werkzeug.utils import secure_filename
    
        # 安全检查
        filename = secure_filename(filename)
        image_path = os.path.join(app.root_path, 'static', 'images', filename)
    
        if os.path.exists(image_path):
            try:
                # 检查是否有题目在使用这个图片
                conn = get_db_connection()
                cursor = conn.cursor(dictionary=True)
                cursor.execute("SELECT id, template_name FROM problem_templates WHERE image_filename = %s", (filename,))
                using_templates = cursor.fetchall()
                cursor.close()
                conn.close()
    
                if using_templates:
                    template_names = [t['template_name'] for t in using_templates]
                    flash(f'无法删除图片，以下题目正在使用：{", ".join(template_names)}', 'danger')
                else:
                    os.remove(image_path)
                    flash('图片删除成功！', 'success')
            except Exception as e:
                print(f"删除图片失败: {str(e)}")
                flash(f'删除图片失败: {str(e)}', 'danger')
        else:
            flash('图片不存在', 'danger')
    
        return redirect(url_for('admin_image_manager'))

    return bp




